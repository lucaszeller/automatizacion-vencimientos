"""Generacion del Excel mensual de control."""

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.utils import MONTH_NAMES


REPORT_COLUMNS = [
    ("Fecha", "fecha_service_actual"),
    ("Nº Soporte", "soporte_actual"),
    ("HORAS ", "horas_actuales"),
    ("TIPO DE SERV.", "tipo_service_actual"),
    ("Cliente", "cliente"),
    ("Serializado", "serial"),
    ("Lugar", "lugar"),
    ("Telefono", "telefono_cliente"),
    ("FECHA ANTE ULTIMO SERV.", "fecha_service_anterior"),
    ("Nº SOPORTE", "soporte_anterior"),
    ("HORAS", "horas_anteriores"),
    ("TIPO DE SERV.", "tipo_service_anterior"),
    ("PROMEDIO HS DE USO.", "promedio_horas_dia"),
    ("FECHA ESTIMADA PROXIMO SERVI.", "fecha_estimada_proximo_service"),
    ("ESTADO", "estado"),
]


def generate_report(df: pd.DataFrame, output_dir: Path, run_date: date | None = None, period_label: str | None = None) -> Path:
    run_date = run_date or date.today()
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"services_control_{run_date.year}.xlsx"
    export_columns = []
    for _, column in REPORT_COLUMNS:
        if column in df.columns:
            export_columns.append(df[column])
        else:
            export_columns.append(pd.Series(index=df.index, dtype="object"))
    export = pd.concat(export_columns, axis=1)
    export.columns = [title for title, _ in REPORT_COLUMNS]

    sheet_name = MONTH_NAMES[run_date.month]
    writer_options = {"engine": "openpyxl"}
    if output_path.exists():
        writer_options.update({"mode": "a", "if_sheet_exists": "replace"})
    with pd.ExcelWriter(output_path, **writer_options) as writer:
        export.to_excel(writer, index=False, sheet_name=sheet_name, startrow=2)
        ws = writer.book[sheet_name]
        _style_sheet(ws)
    return output_path


def _style_sheet(ws) -> None:
    max_col = ws.max_column
    max_row = ws.max_row

    # ── Estilos definidos una sola vez ────────────────────────────────────────
    current_fill  = PatternFill("solid", fgColor="C6E0B4")
    previous_fill = PatternFill("solid", fgColor="F8CBAD")
    next_fill     = PatternFill("solid", fgColor="70AD47")
    status_fill   = PatternFill("solid", fgColor="A9D18E")
    alert_fill    = PatternFill("solid", fgColor="F4CCCC")
    warning_fill  = PatternFill("solid", fgColor="FFF2CC")
    error_fill    = PatternFill("solid", fgColor="EADCF8")
    no_fill       = PatternFill()

    thin   = Side(style="thin",   color="000000")
    medium = Side(style="medium", color="7F7F7F")
    border         = Border(left=thin,   right=thin,   top=thin,   bottom=thin)
    section_border = Border(left=medium, right=medium, top=medium, bottom=medium)

    font_header  = Font(bold=True, size=14)
    font_col_hdr = Font(bold=True, size=11)
    font_next    = Font(bold=True, size=12)
    font_alert   = Font(bold=True, color="9C0006")
    font_warning = Font(bold=True, color="7F6000")
    font_error   = Font(bold=True, color="7030A0")

    align_center_cont = Alignment(horizontal="centerContinuous", vertical="center")
    align_center      = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_top         = Alignment(vertical="top", wrap_text=True)

    current_start, current_end   = 1, 8
    previous_start, previous_end = 9, 13
    next_col, state_col          = 14, 15

    # Fill por columna (para filas de datos y cabecera)
    col_fills: dict[int, PatternFill] = {}
    for c in range(current_start,  current_end  + 1): col_fills[c] = current_fill
    for c in range(previous_start, previous_end + 1): col_fills[c] = previous_fill
    col_fills[next_col]  = next_fill
    col_fills[state_col] = status_fill

    # Fill y font del estado (en columna 15)
    _state_fill = {
        "PROXIMO_A_VENCER": alert_fill,   "VENCE_HOY": alert_fill,
        "VENCIDO":          alert_fill,   "SIN_DATOS": warning_fill,
        "SIN_MAIL":         warning_fill, "SIN_CONTACTO": warning_fill,
        "ERROR_CALCULO":    error_fill,
    }
    _state_font = {
        "PROXIMO_A_VENCER": font_alert,  "VENCE_HOY": font_alert,
        "VENCIDO":          font_alert,  "SIN_DATOS": font_warning,
        "SIN_MAIL":         font_warning,"SIN_CONTACTO": font_warning,
        "ERROR_CALCULO":    font_error,
    }

    DATE_COLS  = {1, previous_start, next_col}
    RATIO_COLS = {13}

    # ── Fila 1: solo borde ───────────────────────────────────────────────────
    for col in range(1, max_col + 1):
        ws.cell(row=1, column=col).border = border

    # ── Fila 2: encabezados de sección ───────────────────────────────────────
    sections = {
        "SERVICIO ACTUAL":   (current_start,  current_end,  current_fill),
        "SERVICIO ANTERIOR": (previous_start, previous_end, previous_fill),
    }
    for title, (start, end, fill) in sections.items():
        for col in range(start, end + 1):
            cell = ws.cell(row=2, column=col)
            cell.fill      = fill
            cell.border    = section_border
            cell.alignment = align_center_cont
        ws.cell(row=2, column=start).value = title
        ws.cell(row=2, column=start).font  = font_header

    for col, value, fill in [
        (next_col,  "PROXIMO SERVICIO", next_fill),
        (state_col, None,               status_fill),
    ]:
        cell = ws.cell(row=2, column=col)
        if value:
            cell.value = value
            cell.font  = font_header
        cell.fill      = fill
        cell.border    = section_border
        cell.alignment = align_center
    ws.row_dimensions[2].height = 28

    # ── Fila 3: cabecera de columnas ──────────────────────────────────────────
    for cell in ws[3]:
        cell.font      = font_col_hdr
        cell.fill      = col_fills.get(cell.column, no_fill)
        cell.border    = border
        cell.alignment = align_center
    ws.row_dimensions[3].height = 28

    # ── Filas de datos: UNA SOLA PASADA ──────────────────────────────────────
    for row in range(4, max_row + 1):
        state = ws.cell(row=row, column=state_col).value
        ws.row_dimensions[row].height = 34

        for col in range(1, max_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.border    = border
            cell.alignment = align_top

            # Fill
            if col == state_col:
                cell.fill = _state_fill.get(state, status_fill)
            else:
                cell.fill = col_fills.get(col, no_fill)

            # Font
            if col == state_col and state in _state_font:
                cell.font = _state_font[state]
            elif col == next_col:
                cell.font = font_next

            # Formato numérico
            if col in DATE_COLS:
                cell.number_format = "dd/mm/yyyy"
            elif col in RATIO_COLS:
                cell.number_format = "0.00"

    # ── Anchos de columna ─────────────────────────────────────────────────────
    widths = {1:13, 2:9.5, 3:7, 4:12, 5:28, 6:23, 7:8, 8:16,
              9:13, 10:13, 11:13, 12:13, 13:13, 14:21, 15:18}
    for col in range(1, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 16)

    ws.auto_filter.ref = f"A3:{get_column_letter(max_col)}{max_row}"
    ws.freeze_panes   = None
