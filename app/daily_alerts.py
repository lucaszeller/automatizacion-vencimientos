"""Revision diaria de services proximos desde la base local."""

from datetime import date, timedelta
from pathlib import Path
import random
import time

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.config import Settings
from app.notifier_whatsapp import build_service_message, normalize_phone, send_whatsapp_text
from app.service_store import get_pending_alerts, record_notice


NEXT_SERVICE_COLUMN = "FECHA ESTIMADA PROXIMO SERVI."
OUTPUT_COLUMNS = [
    "Cliente",
    "Serializado",
    "Lugar",
    "Fecha",
    "Nº Soporte",
    "HORAS ",
    "TIPO DE SERV.",
    NEXT_SERVICE_COLUMN,
    "Dias restantes",
    "ESTADO",
]


def run_daily_alert_check(settings: Settings, logger, today: date | None = None) -> Path:
    today = today or date.today()
    logger.info("Inicio de revision diaria de vencimientos")
    logger.info("Base local revisada: %s", settings.services_db_path)
    upcoming = find_upcoming_services(settings, today=today)
    output_path = settings.output_dir / f"services_proximos_{today.isoformat()}.xlsx"
    write_daily_alert_report(upcoming, output_path)
    stats = send_daily_whatsapp_alerts(upcoming, settings, logger)
    removed = cleanup_old_daily_reports(settings, today)
    logger.info("Services proximos detectados: %s", len(upcoming))
    logger.info("WhatsApp diarios enviados: %s", stats["enviados"])
    logger.info("WhatsApp diarios simulados: %s", stats["simulados"])
    logger.info("Errores WhatsApp diarios: %s", stats["errores"])
    logger.info("Reporte diario generado: %s", output_path)
    if removed:
        logger.info("Reportes diarios viejos eliminados: %s", removed)
    logger.info("Fin de revision diaria de vencimientos")
    return output_path


def find_upcoming_services(settings: Settings, today: date | None = None) -> pd.DataFrame:
    return get_pending_alerts(settings, today=today)


def cleanup_old_daily_reports(settings: Settings, today: date) -> int:
    """Elimina archivos services_proximos_*.xlsx mas viejos que dias_conservar_reportes_diarios."""
    keep_days = settings.dias_conservar_reportes_diarios
    cutoff = today - timedelta(days=keep_days)
    removed = 0
    for path in settings.output_dir.glob("services_proximos_*.xlsx"):
        date_str = path.stem.replace("services_proximos_", "")
        try:
            file_date = date.fromisoformat(date_str)
        except ValueError:
            continue
        if file_date < cutoff:
            try:
                path.unlink()
                removed += 1
            except OSError:
                pass
    return removed


def send_daily_whatsapp_alerts(df: pd.DataFrame, settings: Settings, logger) -> dict:
    stats = {"enviados": 0, "simulados": 0, "errores": 0}
    if df.empty:
        return stats

    rows = list(df.iterrows())
    for i, (_, row) in enumerate(rows):
        destination = normalize_phone(row["telefono_cliente"], settings.whatsapp_country_prefix, logger)
        result = "SIMULADO" if settings.modo_prueba else "ENVIADO"
        error = ""
        try:
            if settings.modo_prueba:
                logger.info("WhatsApp diario simulado: %s - %s", destination, row["serial"])
                stats["simulados"] += 1
            else:
                # Pausa aleatoria antes de cada envio (excepto el primero) para evitar ban
                if i > 0:
                    from app.notifier_whatsapp import dynamic_delay
                    dynamic_delay(settings.whatsapp_delay_min, settings.whatsapp_delay_max, logger)
                send_whatsapp_text(destination, build_service_message(row), settings, logger)
                logger.info("WhatsApp diario enviado: %s - %s", destination, row["serial"])
                stats["enviados"] += 1
            record_notice(settings, row, result, error)
        except Exception as exc:
            error = str(exc)
            stats["errores"] += 1
            record_notice(settings, row, "ERROR", error)
            logger.exception("ERROR WHATSAPP DIARIO: %s - %s", row.get("telefono_cliente"), error)
    return stats


def write_daily_alert_report(df: pd.DataFrame, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    export = _export_daily_rows(df)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        export.to_excel(writer, index=False, sheet_name="PROXIMOS")
        ws = writer.book["PROXIMOS"]
        _style_daily_sheet(ws)


def _export_daily_rows(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return pd.DataFrame(columns=OUTPUT_COLUMNS)
    export = pd.DataFrame(
        {
            "Cliente": df["cliente"],
            "Serializado": df["serial"],
            "Lugar": df["lugar"],
            "Fecha": df["fecha_service_actual"],
            "Nº Soporte": df["soporte_actual"],
            "HORAS ": df["horas_actuales"],
            "TIPO DE SERV.": df["tipo_service_actual"],
            NEXT_SERVICE_COLUMN: df["fecha_estimada_proximo_service"],
            "Dias restantes": df["dias_restantes"],
            "ESTADO": df["estado"],
        }
    )
    return export[OUTPUT_COLUMNS]


def _style_daily_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="70AD47")
    alert_fill = PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if ws.cell(row=row, column=9).value in (0, "0"):
                cell.fill = alert_fill

    widths = {
        1: 30, 2: 24, 3: 10, 4: 13, 5: 11,
        6: 9, 7: 14, 8: 22, 9: 14, 10: 20,
    }
    for col in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 16)
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=4).number_format = "dd/mm/yyyy"
        ws.cell(row=row, column=8).number_format = "dd/mm/yyyy"
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
    ws.freeze_panes = "A2"


def send_test_alert(settings, logger) -> str:
    """Toma el primer vencimiento pendiente y envia el mensaje al numero de prueba (WHATSAPP_TEST_NUMBER).
    Retorna un string describiendo el resultado.
    """
    from app.notifier_whatsapp import build_service_message, normalize_phone, send_whatsapp_text

    if not settings.whatsapp_test_number:
        raise RuntimeError(
            "WHATSAPP_TEST_NUMBER no esta configurado en .env. "
            "Agregue su numero personal para usar esta funcion."
        )

    upcoming = find_upcoming_services(settings)
    test_dest = normalize_phone(settings.whatsapp_test_number, settings.whatsapp_country_prefix, logger)

    if upcoming.empty:
        # Sin vencimientos reales: enviar mensaje de muestra para verificar la conexion
        message = (
            "Hola.\n\n"
            "Este es un mensaje de prueba del sistema de alertas de services.\n\n"
            "Modelo/equipo: Autoelevador Ejemplo\n"
            "Serial: SER-TEST-001\n"
            "Horas actuales: 1500\n"
            "Fecha estimada de vencimiento: (sin vencimientos proximos en la base)\n\n"
            "La conexion con WhatsApp funciona correctamente.\n"
            "Servicio Tecnico."
        )
        logger.info("Sin vencimientos proximos. Enviando mensaje de muestra a %s.", test_dest)
        send_whatsapp_text(test_dest, message, settings, logger)
        logger.info("Mensaje de prueba enviado correctamente a %s", test_dest)
        return f"Mensaje de muestra enviado a {test_dest}.\n(No habia vencimientos proximos en la base local.)"

    row = upcoming.iloc[0]
    message = build_service_message(row)

    logger.info("Enviando mensaje de prueba a %s (en lugar de %s)", test_dest, row.get("telefono_cliente", ""))
    send_whatsapp_text(test_dest, message, settings, logger)
    logger.info("Mensaje de prueba enviado correctamente a %s", test_dest)

    return (
        f"Mensaje enviado a {test_dest}.\n"
        f"Cliente original: {row.get('cliente', '')}\n"
        f"Equipo: {row.get('equipo', '')} | Serial: {row.get('serial', '')}\n"
        f"Vencimiento: {row.get('fecha_estimada_proximo_service', '')}"
    )
